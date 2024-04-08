import os
import re
from pdfminer.high_level import extract_text
from openpyxl import Workbook
from flask import Flask, request, send_file
from docx import Document
import comtypes.client
import comtypes


app = Flask(__name__)

def extract_emails_and_phones(text):
    emails = re.findall(r'[\w\.-]+@[\w\.-]+', text)
    phones = re.findall(r'\b\d{3}[-.\s]?\d{3}[-.\s]?\d{4}\b', text)
    return emails, phones

def parse_resume(file_path):
    text=""

    if file_path.endswith('.pdf'):
        text = extract_text(file_path)
        text = re.sub(r'(?<! ) (?! )', '', text)
        text=re.sub('\t'," ",text)
        text = text.replace('\n', ' ').strip()
    elif file_path.endswith('.docx'):
        doc = Document(file_path)
        text = " ".join([paragraph.text for paragraph in doc.paragraphs])
    elif file_path.endswith('.doc'):
        try:
            comtypes.CoInitialize()
            word_app = comtypes.client.CreateObject("Word.Application")
            doc = word_app.Documents.Open(file_path)
            text = doc.Content.Text
            doc.Close()
            word_app.Quit()
            illegal_chars_pattern = r'[^\x20-\x7E]' 
            text = re.sub(illegal_chars_pattern, '', text)
        except Exception as e:
            print(f"error printed in doc : {e} {file_path}")
    else:
        # Unsupported file format
        return "",[], []

    emails, phones = extract_emails_and_phones(text)
    return text,emails, phones

@app.route('/')
def index():
    return '''
    <html>
    <head>
            <style>
            body {
                font-family: Arial, sans-serif;
                background-color: #f2f2f2;
                margin: 0;
                padding: 0;
            }

            h1 {
                color: #333;
                text-align: center;
                margin-top: 2.3rem;
            }

            form {
                width: 50%;
                margin: 0 auto;
                padding: 20px;
                background-color: #fff;
                border-radius: 8px;
                box-shadow: 0 4px 8px rgba(0, 0, 0, 0.1);
            }

            input[type="file"] {
                display: block;
                margin: 10px 0;
                border: 2px solid #4CAF50;
                background-color: rgb(183, 225, 183);
                border-radius: 4px;
                padding: 10px;
                width: calc(100% - 100px); 
                height: 45px;
                box-sizing: border-box; 
            }

            input[type="file"]:hover {
                border-color: #45a049;
            }

            input[type="submit"] {
                background-color: #4CAF50;
                color: white;
                padding: 10px 20px;
                border: none;
                border-radius: 4px;
                cursor: pointer;
                font-size: 16px;
            }

            input[type="submit"]:hover {
                background-color: #45a049;
            }
        </style>
    </head>
        <body>
            <h1>Upload Resume</h1>
            <form method="post" action="/upload" enctype="multipart/form-data">
                <input type="file" name="file" multiple>
                <input type="submit" value="Upload">
            </form>
        </body>
    </html>
    '''

@app.route('/upload', methods=['POST'])
def upload_file():
    uploaded_files = request.files.getlist("file")
    if not uploaded_files:
        return "No files selected"
    text_list=[]
    emails_list = []
    phones_list = []

    for file in uploaded_files:
        # Create the "uploads" directory if it doesn't exist
        if not os.path.exists('uploads'):
            os.makedirs('uploads')
        # Save the uploaded file temporarily
        file_path = os.path.join('uploads', file.filename)
        file.save(file_path)
        text,emails, phones = parse_resume(file_path)
        illegal_chars_pattern = r'[^\x20-\x7E]'
        text = re.sub(illegal_chars_pattern, '', text)
        text_list.append(text)
        emails_list.extend(emails)
        phones_list.extend(phones)
        # Remove the temporary file after processing
        os.remove(file_path)

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Text','Email', 'Phone'])
    for text,email, phone in zip(text_list,emails_list, phones_list):
        sheet.append([text,email, phone])
    excel_file = 'resumes_data.xlsx'
    workbook.save(excel_file)

    return send_file(excel_file, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)